"""Business-case workbook with live formulas: change an assumption and every result updates in Excel.

Sheets: Inputs (measured values and assumptions), Options (costs, break-even, scenario), Sensitivity
(break-even uplift for option B by margin and discount), Pilot (detectable uplift, sample size).
tests/test_workbook.py recalculates the workbook with the `formulas` engine and checks it against model.py.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from winback.model import ASSUMPTIONS, Inputs

HEAD = PatternFill("solid", fgColor="1F3864")
EDIT = PatternFill("solid", fgColor="FFF2CC")


def _head(ws, row: int) -> None:
    for c in ws[row]:
        if c.value is not None:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD


def build(i: Inputs, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws.append(["Input", "Value", "Unit", "Source"])
    _head(ws, 1)
    rows = [
        ("At-risk customers targeted", i.n, "customers", "Measured: RFM segment on 2011-12-01"),
        ("Return rate without programme (90 days)", i.p0, "share", "Measured: same-season snapshot 2010-12-01"),
        ("Revenue per returning customer (90 days)", i.revenue_per_returner, "GBP", "Measured: same snapshot"),
        ("Top-value customers to call", i.top_value_customers, "customers", "Measured: top 20 % by prior spend"),
        ("Gross margin", ASSUMPTIONS["gross_margin"]["value"], "share", "ASSUMPTION: " + ASSUMPTIONS["gross_margin"]["basis"]),
        ("Discount, option B", ASSUMPTIONS["discount"]["value"], "share", "ASSUMPTION: " + ASSUMPTIONS["discount"]["basis"]),
        ("Set-up effort", ASSUMPTIONS["setup_days"]["value"], "days", "ASSUMPTION: " + ASSUMPTIONS["setup_days"]["basis"]),
        ("Analyst day rate", ASSUMPTIONS["day_rate_gbp"]["value"], "GBP/day", "ASSUMPTION: " + ASSUMPTIONS["day_rate_gbp"]["basis"]),
        ("Minutes per call", ASSUMPTIONS["call_minutes"]["value"], "minutes", "ASSUMPTION: " + ASSUMPTIONS["call_minutes"]["basis"]),
        ("Account manager hourly cost", ASSUMPTIONS["call_rate_gbp_per_hour"]["value"], "GBP/hour", "ASSUMPTION: " + ASSUMPTIONS["call_rate_gbp_per_hour"]["basis"]),
    ]
    for r in rows:
        ws.append(list(r))
    for r in range(6, 12):
        ws.cell(row=r, column=2).fill = EDIT
    for r in (3, 6, 7):
        ws.cell(row=r, column=2).number_format = "0.0%"
    ws.cell(row=4, column=2).number_format = "#,##0.00"
    ws.append([])
    ws.append(["Yellow cells are assumptions: change them and every sheet recalculates."])
    for col, w in zip("ABCD", (42, 12, 11, 70)):
        ws.column_dimensions[col].width = w
    # Named shortcuts used below (cell addresses on Inputs).
    N, P0, R, TOP, M, D, DAYS, RATE, MINS, HRATE = (f"Inputs!$B${r}" for r in range(2, 12))

    ws = wb.create_sheet("Options")
    ws.append(["", "A. Do nothing", "B. Email + discount", "C. Reminder + calls"])
    _head(ws, 1)
    ws.append(["Fixed cost (GBP)", 0, f"={DAYS}*{RATE}", f"={DAYS}*{RATE}+{TOP}*{MINS}/60*{HRATE}"])
    ws.append(["Break-even return rate", f"={P0}", f"=({P0}*{M}+C2/({N}*{R}))/({M}-{D})", f"={P0}+D2/({N}*{R}*{M})"])
    ws.append(["Uplift needed to break even (points)", 0, f"=(C3-{P0})*100", f"=(D3-{P0})*100"])
    ws.append([])
    ws.append(["Scenario: assumed uplift (points)", 3])
    ws.cell(row=6, column=2).fill = EDIT
    ws.append(["Return rate in scenario", f"={P0}", f"={P0}+$B$6/100", f"={P0}+$B$6/100"])
    ws.append(["Gross profit in scenario (GBP)", f"={N}*B7*{R}*{M}", f"={N}*C7*{R}*({M}-{D})-C2", f"={N}*D7*{R}*{M}-D2"])
    ws.append(["Gain vs doing nothing (GBP)", 0, "=C8-$B$8", "=D8-$B$8"])
    ws.append([])
    ws.append(["Option C: gain per extra point of return rate (GBP)", f"={N}*0.01*{R}*{M}"])
    ws.append(["Option C: worst case, no uplift at all (GBP)", "=-D2"])
    for r, fmt in [(2, "#,##0"), (3, "0.0%"), (4, "0.0"), (7, "0.0%"), (8, "#,##0"), (9, "#,##0")]:
        for c in "BCD":
            ws[f"{c}{r}"].number_format = fmt
    ws["B11"].number_format = "#,##0"
    ws["B12"].number_format = "#,##0"
    for col, w in zip("ABCD", (50, 16, 20, 20)):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Sensitivity")
    margins = [0.25, 0.30, 0.35, 0.40, 0.45, 0.50]
    discounts = [0.05, 0.10, 0.15]
    ws.append(["Option B: uplift needed to break even (points)", "", "", "", "", "", ""])
    ws.append(["Discount \\ gross margin"] + margins)
    _head(ws, 2)
    for j, dsc in enumerate(discounts, start=3):
        row = [dsc]
        for k in range(len(margins)):
            col = chr(ord("B") + k)
            row.append(f"=IF({col}$2<=$A{j},\"never\",((({P0})*{col}$2+Options!$C$2/({N}*{R}))/({col}$2-$A{j})-{P0})*100)")
        ws.append(row)
    ws.append([])
    ws.append(["Option C: uplift needed to break even (points)"] + [f"=Options!$D$2/({N}*{R}*{c})*100" for c in margins])
    for r in ws.iter_rows(min_row=2, max_row=2, min_col=2, max_col=7):
        for c in r:
            c.number_format = "0%"
    for r in range(3, 6):
        ws.cell(row=r, column=1).number_format = "0%"
        for c in range(2, 8):
            ws.cell(row=r, column=c).number_format = "0.0"
    for c in range(2, 8):
        ws.cell(row=7, column=c).number_format = "0.00"
    ws.column_dimensions["A"].width = 46

    ws = wb.create_sheet("Pilot")
    ws.append(["Pilot design", "Value", "Note"])
    _head(ws, 1)
    ws.append(["Customers per group (half the segment)", f"=INT({N}/2)", ""])
    ws.append(["Significance level (two-sided)", 0.05, ""])
    ws.append(["Power", 0.8, ""])
    ws.append(["Smallest uplift detectable (points)", f"=(NORM.S.INV(1-B3/2)+NORM.S.INV(B4))*SQRT(2*{P0}*(1-{P0})/B2)*100",
               "Normal approximation; model.py solves it exactly"])
    ws.append(["Customers per group to detect option C's break-even", f"=ROUNDUP((NORM.S.INV(1-B3/2)+NORM.S.INV(B4))^2*2*{P0}*(1-{P0})/(Options!D4/100)^2,0)",
               "Same approximation"])
    for r, fmt in [(3, "0%"), (4, "0%"), (5, "0.0"), (6, "#,##0")]:
        ws.cell(row=r, column=2).number_format = fmt
    for col, w in zip("ABC", (52, 12, 48)):
        ws.column_dimensions[col].width = w
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical="center", wrap_text=isinstance(c.value, str) and len(c.value) > 60)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
